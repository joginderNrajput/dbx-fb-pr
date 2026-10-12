"""Information: how to read the workbook, what was scanned, the limits."""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from dbx_premig.excel import theme
from dbx_premig.excel.layout import back_link, branding_band, nav_link, set_widths
from dbx_premig.excel.writer import DATA_SHEETS
from dbx_premig.models import DiscoveryReport
from dbx_premig.version import APP_NAME, APP_VERSION

__all__ = ["LIMITS", "TITLE", "write_information"]

TITLE = "Information  ·  how to read this workbook"

LIMITS = (
    "Discovery only reads. Live mode issues GET requests (and, for service-principal auth, one "
    "OAuth token request); it never creates, changes or runs anything in the workspace.",
    "What is visible depends on the identity used: objects the token cannot see are missing, "
    "and endpoints it may not call are listed on Warnings. Use a workspace admin for full "
    "coverage.",
    "Code features come from pattern matching on notebook source, not from running it. "
    "Dynamic paths and notebook names built at run time are not resolved.",
    "Secret values, connection options and credentials are never read; only names and counts.",
    "Job-run history, notebook / job / cluster ACLs, Genie spaces, vector search indexes and "
    "DBFS root files are not inventoried.",
    "Readiness scores, verdicts and effort points are heuristics of this tool (tunable in "
    "features.yaml and the settings), not an industry standard or an hours estimate.",
)

_INTRO = (
    "A discovery report: an inventory of everything found in the Databricks workspace (or the "
    "offline export), with a first readiness classification for Microsoft Fabric. It is the "
    "input to scoping; nothing here is an estimate of hours."
)
_NEXT = (
    "Review Warnings first and fill any gaps (for example re-run with an admin token). Then "
    "use Assets and Findings to scope the Databricks -> Fabric migration: notebooks become "
    "Fabric notebooks, jobs become Data Factory pipelines, Unity Catalog tables become "
    "Lakehouse tables or shortcuts, SQL warehouses map to Fabric Warehouse, dashboards are "
    "rebuilt in Power BI."
)

_EXTRA_SHEETS = (
    (
        "Discovery",
        "The dashboard: KPI tiles, scan summary and charts. Every figure links to Totals.",
    ),
    ("Information", "How to read this workbook."),
    (
        "Totals",
        "The single source of truth: live totals, readiness by type, findings by rule "
        "and the editable effort weights.",
    ),
)


def _section(ws: Worksheet, row: int, text: str) -> None:
    ws.cell(row=row, column=1, value=text).font = theme.SECTION_FONT


def _para(ws: Worksheet, row: int, text: str, height: float = 56) -> None:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = theme.BODY_FONT
    cell.alignment = theme.WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.row_dimensions[row].height = height


def write_information(ws: Worksheet, report: DiscoveryReport) -> None:
    set_widths(ws, [30, 22, 22, 22, 22, 44])
    branding_band(ws, TITLE, 6)
    back_link(ws, "How to read this workbook.")
    _section(ws, 8, "What this workbook is")
    _para(ws, 9, _INTRO)
    _section(ws, 11, "What happens next")
    _para(ws, 12, _NEXT, 70)

    _section(ws, 14, "What each sheet holds")
    row = 15
    for i, h in enumerate(("Sheet", "Contents"), start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = theme.HEADER_FONT
        cell.fill = theme.HEADER_FILL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    entries = [
        _EXTRA_SHEETS[0],
        *((s.name, s.description) for s in DATA_SHEETS),
        *_EXTRA_SHEETS[1:],
    ]
    for name, text in entries:
        row += 1
        link = ws.cell(row=row, column=1, value=nav_link(name))
        link.font = theme.font(theme.GREEN, 10, True)
        desc = ws.cell(row=row, column=2, value=text)
        desc.font = theme.BODY_FONT
        desc.alignment = theme.WRAP
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 30

    row += 2
    _section(ws, row, "What was scanned")
    row += 1
    for i, h in enumerate(("Property", "Value"), start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = theme.HEADER_FONT
        cell.fill = theme.HEADER_FILL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    scan = report.scan
    props = (
        (
            "Mode",
            "Live workspace (REST API, read-only)" if scan.mode == "live" else "Offline export",
        ),
        ("Workspace / input", scan.source),
        ("Cloud", scan.cloud or "-"),
        ("Workspace id", scan.workspace_id or "-"),
        ("Unity Catalog metastore", scan.metastore or "-"),
        ("Collectors run", ", ".join(scan.collectors) or "-"),
        ("Captured", scan.captured_at),
        ("Workbook generated", scan.generated_at),
        ("Tool", f"{APP_NAME} {APP_VERSION}"),
    )
    for label, value in props:
        row += 1
        ws.cell(row=row, column=1, value=label).font = theme.font(theme.INK, 10, True)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = theme.BODY_FONT
        cell.data_type = "s"
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)

    row += 2
    _section(ws, row, "Limits of this discovery")
    for text in LIMITS:
        row += 1
        _para(ws, row, f"•  {text}", 30)
