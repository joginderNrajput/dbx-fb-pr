"""Dashboard charts and their hidden helper sheets (``=Totals!...`` formulas only)."""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from dbx_premig.excel import theme

if TYPE_CHECKING:
    from dbx_premig.excel.sheets.totals import TotalsRefs

__all__ = ["CHART_ANCHORS", "add_dashboard_charts", "write_chart_helpers"]

#: Top-left cells of the three dashboard charts (kept apart so they never overlap).
CHART_ANCHORS = ("H30", "Z30", "AR30")
_TOP_TYPES = 12
_TOP_FINDINGS = 10
_SIZE = (12.6, 8.5)  # cm: 17 dashboard columns, so the charts never overlap


def _helper(wb: Workbook, name: str, rows: list[tuple[str, str]]) -> Worksheet:
    ws = wb.create_sheet(name)
    ws["A1"], ws["B1"] = "Label", "Value"
    for i, (label_ref, value_ref) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=label_ref)
        ws.cell(row=i, column=2, value=value_ref)
    ws.sheet_state = "hidden"
    return ws


def write_chart_helpers(wb: Workbook, refs: TotalsRefs) -> None:
    _helper(wb, "ChartReadiness", [(level, f"={ref}") for level, ref in refs.readiness])
    _helper(
        wb,
        "ChartAssetTypes",
        [(f"={label}", f"={count}") for label, count, _ in refs.asset_types[:_TOP_TYPES]],
    )
    _helper(
        wb,
        "ChartTopFindings",
        [(f"={label}", f"={count}") for label, count in refs.top_findings[:_TOP_FINDINGS]],
    )


def _labels() -> DataLabelList:
    labels = DataLabelList()
    labels.showVal = True
    labels.showCatName = False
    labels.showSerName = False
    labels.showPercent = False
    labels.showLegendKey = False
    return labels


def _rows(n: int) -> int:
    return max(n, 1) + 1


def add_dashboard_charts(ws: Worksheet, refs: TotalsRefs) -> None:
    wb = ws.parent
    readiness = wb["ChartReadiness"]
    doughnut = DoughnutChart(holeSize=55)
    doughnut.title = "Assets by Fabric readiness"
    data = Reference(readiness, min_col=2, min_row=1, max_row=_rows(len(refs.readiness)))
    cats = Reference(readiness, min_col=1, min_row=2, max_row=_rows(len(refs.readiness)))
    doughnut.add_data(data, titles_from_data=True)
    doughnut.set_categories(cats)
    series = doughnut.series[0]
    for idx, level in enumerate(("Compatible", "Partial", "Manual", "Unsupported")):
        point = DataPoint(idx=idx)
        point.graphicalProperties = GraphicalProperties(solidFill=theme.READINESS_COLOURS[level][1])
        series.dPt.append(point)
    doughnut.dataLabels = _labels()
    doughnut.legend.position = "b"
    doughnut.width, doughnut.height = _SIZE
    ws.add_chart(doughnut, CHART_ANCHORS[0])

    types = wb["ChartAssetTypes"]
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Assets by type"
    bar.legend = None
    n = len(refs.asset_types[:_TOP_TYPES])
    bar.add_data(Reference(types, min_col=2, min_row=1, max_row=_rows(n)), titles_from_data=True)
    bar.set_categories(Reference(types, min_col=1, min_row=2, max_row=_rows(n)))
    bar.series[0].graphicalProperties = GraphicalProperties(solidFill=theme.GREEN)
    bar.y_axis.majorGridlines = None
    bar.x_axis.scaling.orientation = "maxMin"
    bar.dataLabels = _labels()
    bar.width, bar.height = _SIZE
    ws.add_chart(bar, CHART_ANCHORS[1])

    top = wb["ChartTopFindings"]
    fbar = BarChart()
    fbar.type = "bar"
    fbar.title = "Top findings (assets affected)"
    fbar.legend = None
    m = len(refs.top_findings[:_TOP_FINDINGS])
    fbar.add_data(Reference(top, min_col=2, min_row=1, max_row=_rows(m)), titles_from_data=True)
    fbar.set_categories(Reference(top, min_col=1, min_row=2, max_row=_rows(m)))
    fbar.series[0].graphicalProperties = GraphicalProperties(solidFill=theme.NAVY)
    fbar.y_axis.majorGridlines = None
    fbar.x_axis.scaling.orientation = "maxMin"
    fbar.dataLabels = _labels()
    fbar.width, fbar.height = _SIZE
    ws.add_chart(fbar, CHART_ANCHORS[2])
