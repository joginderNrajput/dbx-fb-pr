"""Write the discovery workbook (the deliverable of the ``discover`` flow)."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook

from dbx_premig.config import Settings
from dbx_premig.errors import ReportWriteError
from dbx_premig.excel.charts import write_chart_helpers
from dbx_premig.excel.dashboard import write_dashboard
from dbx_premig.excel.sheets.information import write_information
from dbx_premig.excel.sheets.totals import write_totals
from dbx_premig.excel.writer import DATA_SHEETS, order_sheets, write_data_sheet
from dbx_premig.logging import get_logger
from dbx_premig.models import DiscoveryReport
from dbx_premig.version import APP_NAME, APP_VERSION

__all__ = ["build_workbook", "save_atomically", "write_discovery_workbook"]

log = get_logger("excel")


def build_workbook(report: DiscoveryReport, settings: Settings) -> Workbook:
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Discovery"
    for spec in DATA_SHEETS:
        ws = wb.create_sheet(spec.name)
        n = write_data_sheet(ws, spec, report, settings)
        log.info("sheet %s: %d rows", spec.name, n)
    write_information(wb.create_sheet("Information"), report)
    refs = write_totals(wb.create_sheet("Totals"), report)
    write_chart_helpers(wb, refs)
    write_dashboard(dashboard, report, refs)
    order_sheets(wb)
    wb.active = 0
    props = wb.properties
    props.title = "Databricks discovery"
    props.subject = "Databricks -> Microsoft Fabric pre-migration discovery"
    props.creator = f"{APP_NAME} {APP_VERSION} (Kanerika)"
    props.keywords = "Databricks, Microsoft Fabric, migration, discovery"
    # Excel recalculates every formula on open (openpyxl stores no cached values).
    wb.calculation.fullCalcOnLoad = True
    return wb


def write_discovery_workbook(report: DiscoveryReport, path: Path, settings: Settings) -> Path:
    """Build and atomically save the workbook. Raises ReportWriteError."""
    return save_atomically(build_workbook(report, settings), path)


def save_atomically(wb: Workbook, path: Path) -> Path:
    """Save via a temporary file in the same folder, then rename. Raises ReportWriteError."""
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
        os.close(fd)
        try:
            wb.save(tmp)
            umask = os.umask(0)
            os.umask(umask)
            os.chmod(tmp, 0o666 & ~umask)  # mkstemp creates 0600; give the usual file mode
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
    except PermissionError as exc:
        raise ReportWriteError(
            f"Cannot write {path.name}: permission denied (is it open in Excel?)."
        ) from exc
    except OSError as exc:
        raise ReportWriteError(f"Cannot write {path.name} ({type(exc).__name__}).") from exc
    return path
