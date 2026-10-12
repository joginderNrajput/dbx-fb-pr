"""Shared worksheet layout: branding band (rows 1-5), header row 6, data from row 7."""

from __future__ import annotations

from collections.abc import Sequence

from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from dbx_premig.excel import theme
from dbx_premig.resources import asset_path

__all__ = [
    "DATA_START_ROW",
    "HEADER_ROW",
    "back_link",
    "branding_band",
    "column_px",
    "data_range",
    "nav_link",
    "title_asset",
]

HEADER_ROW = 6
DATA_START_ROW = 7
#: Formula ranges reach at least this row so a reviewer can append rows by hand.
MIN_RANGE_END = 100_000

_BAND_ROW_HEIGHTS = {1: 24.0, 2: 24.0, 3: 20.0, 4: 10.0, 5: 18.0}
_MIN_BAND_PX = 1500
_LOGO_SIZE = (200, 45)  # Kanerika logo, px
_PLATFORM_LOGO_SIZE = (234, 46)  # Databricks -> Fabric strip, px
_LOGO_X, _LOGO_Y = 14, 15  # px: Kanerika logo position (as in the reference workbook)
_TITLE_X = _LOGO_X + 200 + 36  # px from the left edge: clear of the Kanerika logo


def column_px(ws: Worksheet, index: int) -> int:
    """Approximate rendered width of a 1-based column, in pixels."""
    dim = ws.column_dimensions.get(get_column_letter(index))
    width = dim.width if dim is not None and dim.width else 8.43
    return int(width * 7 + 5)


def data_range(sheet: str, column: str, rows: int) -> str:
    """Absolute range over a data column, e.g. 'Jobs'!$A$7:$A$100000."""
    end = max(MIN_RANGE_END, DATA_START_ROW + rows + 1000)
    return f"'{sheet}'!${column}${DATA_START_ROW}:${column}${end}"


def nav_link(sheet: str, text: str | None = None) -> str:
    label = (text or sheet).replace('"', "'")
    return f'=HYPERLINK("#\'{sheet}\'!A1","{label}")'


def back_link(ws: Worksheet, description: str, home: str = "Discovery") -> None:
    ws["A5"] = nav_link(home, f"«  Back to {home}")
    ws["A5"].font = theme.LINK_FONT
    ws["B5"] = description
    ws["B5"].font = theme.NOTE_FONT


def _image(name: str, size: tuple[int, int]) -> XLImage:
    img = XLImage(str(asset_path(name)))
    img.width, img.height = size
    return img


def title_asset(sheet: str) -> str:
    """File name of the pre-rendered title overlay for a sheet (see scripts/extract_assets.py)."""
    return f"titles/{sheet}.png"


def anchor_at(ws: Worksheet, x_px: int) -> tuple[int, int]:
    """(0-based column, pixel offset inside it) for a horizontal position on the sheet.

    Excel caps a picture's column offset at that column's width, so an offset must never be
    larger than its column: a picture placed "250 px into column A" on a sheet whose column A
    is 130 px wide would otherwise slide left (onto the Kanerika logo).
    """
    col, start = 0, 0
    while True:
        width = column_px(ws, col + 1)
        if x_px < start + width:
            return col, x_px - start
        start += width
        col += 1


def _one_cell(x_px: int, y_px: int, size: tuple[int, int], ws: Worksheet) -> OneCellAnchor:
    col, off = anchor_at(ws, x_px)
    return OneCellAnchor(
        _from=AnchorMarker(col=col, colOff=pixels_to_EMU(off), row=0, rowOff=pixels_to_EMU(y_px)),
        ext=XDRPositiveSize2D(pixels_to_EMU(size[0]), pixels_to_EMU(size[1])),
    )


def _title_overlay(ws: Worksheet, sheet: str, max_right_px: int) -> None:
    """White title text as a transparent image on top of the banner.

    Cell text would sit *under* the banner picture, so the title is pre-rendered at 2x
    and shown at half size, never stretched.
    """
    path = asset_path(title_asset(sheet))
    if not path.is_file():
        return
    img = XLImage(str(path))
    width, height = int(img.width) // 2, int(img.height) // 2
    room = max_right_px - _TITLE_X
    if width > room > 0:  # never run into the platform logo: shrink, keep the aspect ratio
        width, height = room, max(1, height * room // width)
    img.width, img.height = width, height
    # Vertically centred on the Kanerika logo.
    y = _LOGO_Y + (_LOGO_SIZE[1] - height) // 2
    img.anchor = _one_cell(_TITLE_X, y, (width, height), ws)
    ws.add_image(img)


def branding_band(ws: Worksheet, title: str, last_col: int, title_key: str | None = None) -> None:
    """Banner across rows 1-4, Kanerika logo left, platform logo right, title in white."""
    title_key = title_key or ws.title
    for row, height in _BAND_ROW_HEIGHTS.items():
        ws.row_dimensions[row].height = height
    ws.sheet_view.showGridLines = False
    # Print / PDF: landscape, one page wide.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    # Stretch the banner over the used columns, and at least _MIN_BAND_PX wide.
    span, col = 0, 0
    while col < last_col or span < _MIN_BAND_PX:
        col += 1
        span += column_px(ws, col)
    banner = _image("banner.png", (span, 76))
    banner.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=0, colOff=0, row=0, rowOff=0),
        to=AnchorMarker(col=col, colOff=0, row=4, rowOff=0),
    )
    ws.add_image(banner)

    logo = _image("kanerika_logo.png", _LOGO_SIZE)
    logo.anchor = _one_cell(_LOGO_X, _LOGO_Y, _LOGO_SIZE, ws)
    ws.add_image(logo)

    # Platform logo: right-aligned inside the banner (generous margin: renderers differ in
    # how wide a column is).
    platform_x = max(_TITLE_X, span - _PLATFORM_LOGO_SIZE[0] - 48)
    platform = _image("platform_logo.png", _PLATFORM_LOGO_SIZE)
    platform.anchor = _one_cell(platform_x, _LOGO_Y, _PLATFORM_LOGO_SIZE, ws)
    ws.add_image(platform)

    _title_overlay(ws, title_key, platform_x - 24)

    # Title text in a cell too (searchable, read by screen readers; the banner covers it).
    acc, title_col = 0, 1
    while acc < _LOGO_SIZE[0] + 40:
        acc += column_px(ws, title_col)
        title_col += 1
    cell = ws.cell(row=2, column=title_col, value=title)
    cell.font = theme.TITLE_FONT
    cell.alignment = theme.LEFT


def set_widths(ws: Worksheet, widths: Sequence[float]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
