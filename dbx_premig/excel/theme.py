"""The ONLY place raw hex colours may appear (a test enforces it).

Palette taken from Kanerika's reference discovery workbook.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

__all__ = [
    "AMBER",
    "BLUE_INPUT",
    "BODY_FONT",
    "CHART_PALETTE",
    "GREEN",
    "GREEN_DARK",
    "GREEN_TINT",
    "GREY_TEXT",
    "HEADER_FILL",
    "HEADER_FONT",
    "INK",
    "NAVY",
    "PURPLE",
    "READINESS_COLOURS",
    "RED",
    "RED_TINT",
    "SLATE",
    "WHITE",
    "YELLOW_INPUT",
    "ZEBRA",
    "fill",
    "font",
]

NAVY = "1B3139"
SLATE = "2C4650"
INK = "1B3139"
GREEN = "009845"
GREEN_DARK = "00733A"
GREEN_TINT = "E8F5F0"
ZEBRA = "F9F7F4"
GREY_TEXT = "54565A"
RED = "C0392B"
RED_TINT = "FDECEA"
AMBER = "B7791F"
AMBER_TINT = "FDF3E1"
PURPLE = "6B21A8"
PURPLE_TINT = "F3E8FF"
WHITE = "FFFFFF"
YELLOW_INPUT = "FFF2CC"
BLUE_INPUT = "0000FF"
BORDER_GREY = "D9DCDE"

#: Series colours for charts, in order.
CHART_PALETTE = (
    "009845",
    "1B3139",
    "4FA3B5",
    "B7791F",
    "7FC8A0",
    "006580",
    "8FB8C4",
    "C0392B",
    "2C4650",
    "A3A5A8",
    "00733A",
    "C9D3D6",
)

#: Fill / font colour per readiness level.
READINESS_COLOURS = {
    "Compatible": (GREEN_TINT, GREEN_DARK),
    "Partial": (AMBER_TINT, AMBER),
    "Manual": (RED_TINT, RED),
    "Unsupported": (PURPLE_TINT, PURPLE),
}


def fill(colour: str) -> PatternFill:
    return PatternFill("solid", fgColor=colour)


def font(colour: str = INK, size: float = 10, bold: bool = False, name: str = "Calibri") -> Font:
    return Font(name=name, size=size, bold=bold, color=colour)


DISPLAY_FONT = "Aptos Display"
HEADER_FILL = fill(NAVY)
HEADER_FONT = font(WHITE, 9, True, DISPLAY_FONT)
BODY_FONT = font(INK, 10)
TITLE_FONT = font(WHITE, 15, True, DISPLAY_FONT)
SECTION_FONT = font(GREEN, 12, True, DISPLAY_FONT)
LINK_FONT = font(GREEN, 9, True)
NOTE_FONT = font(GREY_TEXT, 9)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
THIN_BOTTOM = Border(bottom=Side(style="thin", color=BORDER_GREY))
