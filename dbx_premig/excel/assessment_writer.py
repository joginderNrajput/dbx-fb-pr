"""Write the assessment workbook (the deliverable of the ``assess`` / ``analyze`` flows)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.worksheet import Worksheet

from dbx_premig.config import Settings
from dbx_premig.excel import theme
from dbx_premig.excel.charts import data_labels
from dbx_premig.excel.dashboard import (
    TILE_COLS,
    TILE_ROWS,
    dashboard_frame,
    summary_header,
    summary_row,
    tile,
)
from dbx_premig.excel.discovery_writer import save_atomically
from dbx_premig.excel.sheets.assessment_report import (
    ASSESSMENT_SHEETS,
    DASHBOARD_TITLE,
    EFFORT_BY_TYPE,
    ESTIMATION,
    SummaryRefs,
    write_about,
    write_rules,
    write_summary,
)
from dbx_premig.excel.writer import order_sheets, write_data_sheet
from dbx_premig.models import AssessmentResult
from dbx_premig.version import APP_NAME, APP_VERSION

__all__ = [
    "ASSESSMENT_CHART_ANCHORS",
    "ASSESSMENT_ORDER",
    "build_assessment_workbook",
    "write_assessment_workbook",
]

_CHARTS = ("ChartComplexity", "ChartAutomation", "ChartHoursByType")
ASSESSMENT_ORDER = (*ASSESSMENT_SHEETS, *_CHARTS)
ASSESSMENT_CHART_ANCHORS = ("H24", "Z24", "AR24")
_SIZE = (12.6, 8.5)
_TOP_TYPES = 10

_TILES = (
    (("total_assets", "ASSETS ASSESSED", "#,##0"), ("simple_count", "SIMPLE", "#,##0"),
     ("medium_count", "MEDIUM", "#,##0"), ("complex_count", "COMPLEX", "#,##0")),
    (("fully_automated", "FULLY AUTOMATED", "#,##0"),
     ("partially_automated", "PARTIALLY AUTOMATED", "#,##0"),
     ("manual", "MANUAL", "#,##0"), ("total_constructs", "CONSTRUCTS", "#,##0")),
    (("total_manual_hours", "MANUAL HOURS", "#,##0"),
     ("total_automation_hours", "HOURS WITH AUTOMATION", "#,##0"),
     ("hours_saved", "HOURS SAVED", "#,##0"), ("saving_pct", "SAVING", "0%")),
)  # fmt: skip


def _helper(wb: Workbook, name: str, rows: list[tuple[str, str]]) -> Worksheet:
    ws = wb.create_sheet(name)
    ws["A1"], ws["B1"] = "Label", "Value"
    for i, (label, value) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
    ws.sheet_state = "hidden"
    return ws


def _doughnut(ws: Worksheet, src: Worksheet, n: int, title: str, colours: tuple[str, ...],
              anchor: str) -> None:  # fmt: skip
    chart = DoughnutChart(holeSize=55)
    chart.title = title
    chart.add_data(Reference(src, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
    chart.set_categories(Reference(src, min_col=1, min_row=2, max_row=n + 1))
    for idx, colour in enumerate(colours):
        point = DataPoint(idx=idx)
        point.graphicalProperties = GraphicalProperties(solidFill=colour)
        chart.series[0].dPt.append(point)
    chart.dataLabels = data_labels()
    chart.legend.position = "b"
    chart.width, chart.height = _SIZE
    ws.add_chart(chart, anchor)


def _write_dashboard(ws: Worksheet, result: AssessmentResult, refs: SummaryRefs, types: int
                     ) -> None:  # fmt: skip
    scan = result.discovery.scan
    comp = result.computation
    dashboard_frame(
        ws,
        DASHBOARD_TITLE,
        f"Discovery: {result.discovery.file_name}   ·   source: {scan.source}   ·   "
        f"assessed {comp.assessment_generated_at}",
        list(ASSESSMENT_SHEETS),
        {"EffortByType": "Effort by type"},
    )
    for tile_row, tiles in zip(TILE_ROWS[:3], _TILES, strict=True):
        for col, (key, label, fmt) in zip(TILE_COLS, tiles, strict=True):
            tile(ws, tile_row, col, refs.ref(key), label, risk=key == "manual", fmt=fmt)

    summary_header(ws, "ASSESSMENT SUMMARY")
    r = refs.ref
    rows: list[tuple[str, object, str, str | None]] = [
        ("Discovery source", scan.source, theme.INK, None),
        ("Discovery mode", "Live workspace" if scan.mode == "live" else "Offline export",
         theme.INK, None),
        ("Discovery generated", scan.generated_at, theme.INK, None),
        ("Assessment generated", comp.assessment_generated_at, theme.INK, None),
        ("Rules", comp.rules_version, theme.INK, None),
        ("Supported constructs", f"={r('supported_constructs')}", theme.GREEN_DARK, "#,##0.0"),
        ("Unsupported constructs", f"={r('unsupported_constructs')}", theme.RED, "#,##0.0"),
        ("Manual development hours", f"={r('manual_dev_hours')}", theme.INK, "#,##0.0"),
        ("Manual testing hours", f"={r('manual_test_hours')}", theme.INK, "#,##0.0"),
        ("Automation development hours", f"={r('automation_dev_hours')}", theme.INK, "#,##0.0"),
        ("Automation testing hours", f"={r('automation_test_hours')}", theme.INK, "#,##0.0"),
        ("Hours saved", f"={r('hours_saved')}", theme.GREEN_DARK, "#,##0.0"),
        ("Saving", f"={r('saving_pct')}", theme.GREEN_DARK, "0.0%"),
        ("Notebooks / jobs / tables", f"={r('count_notebooks')}&\" / \"&{r('count_jobs')}"
         f"&\" / \"&{r('count_tables')}", theme.INK, None),
        ("Assessed by", f"{APP_NAME} {APP_VERSION} (Kanerika)", theme.INK, None),
    ]  # fmt: skip
    for i, (label, value, colour, num_fmt) in enumerate(rows):
        summary_row(ws, 8 + i, label, value, i % 2 == 0, colour, num_fmt)

    ws.cell(row=22, column=8, value="AT A GLANCE").font = theme.font(
        theme.GREEN, 9, True, theme.DISPLAY_FONT
    )
    wb = ws.parent
    _doughnut(
        ws,
        wb["ChartComplexity"],
        3,
        "Assets by complexity",
        (theme.GREEN, theme.AMBER, theme.RED),
        ASSESSMENT_CHART_ANCHORS[0],
    )
    _doughnut(
        ws,
        wb["ChartAutomation"],
        3,
        "Assets by automation",
        (theme.GREEN_DARK, theme.SLATE, theme.RED),
        ASSESSMENT_CHART_ANCHORS[1],
    )
    src = wb["ChartHoursByType"]
    bar = BarChart()
    bar.type = "bar"
    bar.grouping = "clustered"
    bar.title = "Hours by asset type: manual vs automation"
    n = max(types, 1)
    bar.add_data(
        Reference(src, min_col=2, max_col=3, min_row=1, max_row=n + 1), titles_from_data=True
    )
    bar.set_categories(Reference(src, min_col=1, min_row=2, max_row=n + 1))
    bar.series[0].graphicalProperties = GraphicalProperties(solidFill=theme.NAVY)
    bar.series[1].graphicalProperties = GraphicalProperties(solidFill=theme.GREEN)
    bar.y_axis.majorGridlines = None
    bar.x_axis.scaling.orientation = "maxMin"
    bar.legend.position = "b"
    bar.width, bar.height = _SIZE
    ws.add_chart(bar, ASSESSMENT_CHART_ANCHORS[2])
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.tabColor = theme.GREEN


def build_assessment_workbook(
    result: AssessmentResult, settings: Settings, rules_text: str, rules_sha256: str
) -> Workbook:
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Assessment"
    write_data_sheet(wb.create_sheet("Estimation"), ESTIMATION, result, settings, "Assessment")
    types = len(EFFORT_BY_TYPE.rows(result))
    write_data_sheet(wb.create_sheet("EffortByType"), EFFORT_BY_TYPE, result, settings,
                     "Assessment")  # fmt: skip
    refs = write_summary(wb.create_sheet("Summary"), result)
    write_rules(wb.create_sheet("Rules"), result, rules_text, rules_sha256)
    write_about(wb.create_sheet("About"), result)
    _helper(wb, "ChartComplexity", [(lvl, f"={refs.ref(k)}") for lvl, k in
            (("Simple", "simple_count"), ("Medium", "medium_count"),
             ("Complex", "complex_count"))])  # fmt: skip
    _helper(wb, "ChartAutomation", [(lvl, f"={refs.ref(k)}") for lvl, k in
            (("Fully automated", "fully_automated"),
             ("Partially automated", "partially_automated"), ("Manual", "manual"))])  # fmt: skip
    hours = _helper(wb, "ChartHoursByType", [])
    hours["B1"], hours["C1"] = "Manual hours", "Automation hours"
    for i in range(min(types, _TOP_TYPES)):
        r = 7 + i
        hours.cell(row=2 + i, column=1, value=f"='EffortByType'!$A${r}")
        hours.cell(row=2 + i, column=2, value=f"='EffortByType'!$I${r}")
        hours.cell(row=2 + i, column=3, value=f"='EffortByType'!$J${r}")
    _write_dashboard(dashboard, result, refs, min(types, _TOP_TYPES))
    order_sheets(wb, ASSESSMENT_ORDER)
    wb.active = 0
    props = wb.properties
    props.title = "Databricks to Fabric assessment"
    props.creator = f"{APP_NAME} {APP_VERSION} (Kanerika)"
    wb.calculation.fullCalcOnLoad = True
    return wb


def write_assessment_workbook(
    result: AssessmentResult, path: Path, settings: Settings, rules_text: str, rules_sha256: str
) -> Path:
    return save_atomically(build_assessment_workbook(result, settings, rules_text, rules_sha256),
                           path)  # fmt: skip
