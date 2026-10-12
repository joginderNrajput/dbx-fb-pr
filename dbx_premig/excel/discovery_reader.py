"""Read a discovery workbook back for the ``assess`` flow (untrusted input).

The discovery workbook is a file format: the data sheets' headers are a contract
(``writer.DATA_SHEETS``). This reader rebuilds every assessable object from the Assets and
Findings rows (joined on ``Asset key``) and **recomputes every count from the data rows**, so
editing a row changes the assessment and editing a total on the Totals sheet does not.

Safety: the file's zip manifest is validated with the same caps as any archive before
openpyxl opens it; only literal cell values are read (formulas are never evaluated).
"""

from __future__ import annotations

import zipfile
from collections import Counter, defaultdict
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from dbx_premig.config import Settings
from dbx_premig.errors import InvalidWorkbookError
from dbx_premig.excel.layout import HEADER_ROW
from dbx_premig.excel.writer import spec_by_name
from dbx_premig.models import AssetInput, DiscoveryWorkbook, Finding, ScanInfo
from dbx_premig.reader import validate_archive

__all__ = ["COUNT_KEYS", "item_counts", "read_discovery_workbook"]

Row = dict[str, Any]
_REQUIRED = ("Assets", "Findings")


def _num(value: Any) -> float:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, int | float):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _rows(ws: Any, sheet: str) -> list[Row]:
    """Data rows of a sheet as {header: value}; headers must match the sheet spec."""
    expected = [c.header for c in spec_by_name(sheet).columns]
    rows_iter: Iterator[tuple[Any, ...]] = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    header_row = next(rows_iter, ())
    headers = [_text(h) for h in header_row[: len(expected)]]
    missing = [h for h in expected if h not in headers]
    if sheet in _REQUIRED and missing:
        raise InvalidWorkbookError(
            f"Sheet {sheet} is not a discovery sheet (missing column(s): {', '.join(missing)})."
        )
    out: list[Row] = []
    for values in rows_iter:
        if not values or values[0] in (None, ""):
            continue
        out.append({h: values[i] if i < len(values) else None for i, h in enumerate(headers)})
    return out


# --------------------------------------------------------------------------- counts

Sheets = dict[str, list[Row]]


def _count(sheet: str) -> Callable[[Sheets], float]:
    return lambda s: float(len(s.get(sheet, [])))


def _where(sheet: str, column: str, *values: str) -> Callable[[Sheets], float]:
    wanted = {v.casefold() for v in values}
    return lambda s: float(
        sum(1 for r in s.get(sheet, []) if _text(r.get(column)).casefold() in wanted)
    )


def _sum(sheet: str, column: str) -> Callable[[Sheets], float]:
    return lambda s: sum(_num(r.get(column)) for r in s.get(sheet, []))


def _nonblank(sheet: str, column: str) -> Callable[[Sheets], float]:
    return lambda s: float(sum(1 for r in s.get(sheet, []) if _text(r.get(column)).strip()))


def _positive(sheet: str, column: str) -> Callable[[Sheets], float]:
    return lambda s: float(sum(1 for r in s.get(sheet, []) if _num(r.get(column)) > 0))


def _avg(sheet: str, column: str) -> Callable[[Sheets], float]:
    def f(s: Sheets) -> float:
        rows = s.get(sheet, [])
        return round(sum(_num(r.get(column)) for r in rows) / len(rows), 2) if rows else 0.0

    return f


_TABLE_ONLY = ("MANAGED", "EXTERNAL", "STREAMING_TABLE", "FOREIGN")

#: Every item count stored with a discovery (column name of the ``totals`` table -> rule).
COUNT_KEYS: dict[str, Callable[[Sheets], float]] = {
    "notebooks": _count("Notebooks"),
    "notebook_lines": _sum("Notebooks", "Lines"),
    "python_notebooks": _where("Notebooks", "Language", "PYTHON"),
    "sql_notebooks": _where("Notebooks", "Language", "SQL"),
    "scala_notebooks": _where("Notebooks", "Language", "SCALA"),
    "r_notebooks": _where("Notebooks", "Language", "R"),
    "notebooks_not_scanned": _where("Notebooks", "Scanned", "No"),
    "repos": _count("Repos"),
    "dependencies": _count("Dependencies"),
    "unresolved_dependencies": _where("Dependencies", "Resolved", "No"),
    "jobs": _count("Jobs"),
    "job_tasks": _count("JobTasks"),
    "scheduled_jobs": _nonblank("Jobs", "Schedule (Quartz cron)"),
    "continuous_jobs": _where("Jobs", "Continuous", "Yes"),
    "dlt_pipelines": _count("DLTPipelines"),
    "clusters": _count("Clusters"),
    "gpu_clusters": _where("Clusters", "GPU", "Yes"),
    "clusters_with_init_scripts": _positive("Clusters", "Init scripts"),
    "instance_pools": _count("InstancePools"),
    "cluster_policies": _count("ClusterPolicies"),
    "sql_warehouses": _count("SQLWarehouses"),
    "libraries": _count("Libraries"),
    "catalogs": _count("Catalogs"),
    "schemas": _count("Schemas"),
    "tables": _where("Tables", "Table type", *_TABLE_ONLY),
    "managed_tables": _where("Tables", "Table type", "MANAGED"),
    "external_tables": _where("Tables", "Table type", "EXTERNAL"),
    "views": _where("Tables", "Table type", "VIEW"),
    "materialized_views": _where("Tables", "Table type", "MATERIALIZED_VIEW"),
    "streaming_tables": _where("Tables", "Table type", "STREAMING_TABLE"),
    "foreign_tables": _where("Tables", "Table type", "FOREIGN"),
    "hive_metastore_tables": _where("Tables", "Catalog", "hive_metastore"),
    "columns": _count("Columns"),
    "volumes": _count("Volumes"),
    "external_locations": _count("ExternalLocations"),
    "storage_credentials": _count("StorageCredentials"),
    "connections": _count("Connections"),
    "functions": _count("Functions"),
    "grants": _count("Grants"),
    "sql_queries": _count("SQLQueries"),
    "dashboards": _count("Dashboards"),
    "alerts": _count("Alerts"),
    "ml_models": _count("Models"),
    "serving_endpoints": _count("ServingEndpoints"),
    "experiments": _count("Experiments"),
    "secret_scopes": _count("SecretScopes"),
    "secrets": _sum("SecretScopes", "Secrets"),
    "mounts": _count("Mounts"),
    "users": _where("Identities", "Type", "User"),
    "groups": _where("Identities", "Type", "Group"),
    "service_principals": _where("Identities", "Type", "Service principal"),
    "assessable_assets": _count("Assets"),
    "compatible_assets": _where("Assets", "Readiness", "Compatible"),
    "partial_assets": _where("Assets", "Readiness", "Partial"),
    "manual_assets": _where("Assets", "Readiness", "Manual"),
    "unsupported_assets": _where("Assets", "Readiness", "Unsupported"),
    "avg_readiness_score": _avg("Assets", "Readiness score"),
    "findings": _count("Findings"),
    "warnings": _count("Warnings"),
}


def item_counts(sheets: Sheets) -> dict[str, float]:
    return {key: rule(sheets) for key, rule in COUNT_KEYS.items()}


# --------------------------------------------------------------------------- assets


def _identity(key: str) -> str:
    return key.split(":", 1)[1] if ":" in key else key


def _metrics(sheets: Sheets) -> dict[str, dict[str, float]]:
    """Size metrics per asset key, from the detail sheets."""
    out: dict[str, dict[str, float]] = defaultdict(dict)
    for r in sheets.get("Notebooks", []):
        out[f"Notebook:{_text(r.get('Path'))}"] = {
            "lines": _num(r.get("Lines")),
            "cells": _num(r.get("Cells")),
            "notebook_calls": _num(r.get("Notebook calls")),
            "sql_statements": _num(r.get("SQL statements")),
            "magic_commands": _num(r.get("Magic commands")),
            "size_bytes": _num(r.get("Size (bytes)")),
        }
    tasks = Counter(_text(r.get("Job id")) for r in sheets.get("JobTasks", []))
    for r in sheets.get("Jobs", []):
        job_id = _text(r.get("Job id"))
        out[f"Job:{job_id}"] = {"tasks": float(tasks.get(job_id, 0))}
    views = {_text(r.get("Full name")): _num(r.get("Definition (chars)"))
             for r in sheets.get("Views", [])}  # fmt: skip
    for r in sheets.get("Tables", []):
        full = _text(r.get("Full name"))
        metrics = {"columns": _num(r.get("Columns")), "definition_chars": views.get(full, 0.0)}
        for t in ("Table", "View", "Materialized View", "Streaming Table"):
            out[f"{t}:{full}"] = metrics
    for r in sheets.get("Functions", []):
        out[f"Function:{_text(r.get('Full name'))}"] = {
            "definition_chars": _num(r.get("Body (chars)"))
        }
    for r in sheets.get("DLTPipelines", []):
        out[f"DLT Pipeline:{_text(r.get('Pipeline id'))}"] = {"libraries": _num(r.get("Libraries"))}
    for r in sheets.get("SQLQueries", []):
        out[f"SQL Query:{_text(r.get('Query id'))}"] = {"lines": _num(r.get("Lines"))}
    return out


def _scan_info(ws: Any) -> ScanInfo:
    props: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        if row and row[0] is not None and len(row) > 1 and row[1] is not None:
            props[_text(row[0])] = _text(row[1])
    mode = props.get("Mode", "")
    return ScanInfo(
        mode="live" if mode.lower().startswith("live") else "offline",
        source=props.get("Workspace / input", ""),
        cloud="" if props.get("Cloud") == "-" else props.get("Cloud", ""),
        workspace_id="" if props.get("Workspace id") == "-" else props.get("Workspace id", ""),
        metastore=""
        if props.get("Unity Catalog metastore") == "-"
        else props.get("Unity Catalog metastore", ""),
        collectors=tuple(
            c.strip()
            for c in props.get("Collectors run", "").split(",")
            if c.strip() != "-" and c.strip()
        ),
        captured_at=props.get("Captured", ""),
        generated_at=props.get("Workbook generated", ""),
    )


def read_discovery_workbook(path: Path, settings: Settings) -> DiscoveryWorkbook:
    """Rebuild the assessable objects and item counts from a discovery workbook."""
    if path.suffix.lower() != ".xlsx":
        raise InvalidWorkbookError(f"{path.name} is not an .xlsx discovery workbook.")
    try:
        if path.stat().st_size > settings.max_file_bytes:
            raise InvalidWorkbookError(f"{path.name} is larger than the safety cap.")
        with zipfile.ZipFile(path) as zf:
            refusal = validate_archive(zf, path.name, settings)
    except (OSError, zipfile.BadZipFile) as exc:
        raise InvalidWorkbookError(f"{path.name} is not a readable .xlsx file.") from exc
    if refusal:
        raise InvalidWorkbookError(f"{path.name} refused: {refusal}.")
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except (InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise InvalidWorkbookError(f"{path.name} could not be opened as a workbook.") from exc
    try:
        missing = [s for s in (*_REQUIRED, "Information") if s not in wb.sheetnames]
        if missing:
            raise InvalidWorkbookError(
                f"{path.name} is not a discovery workbook (no {', '.join(missing)} sheet)."
            )
        from dbx_premig.excel.writer import DATA_SHEETS

        sheets: Sheets = {
            spec.name: _rows(wb[spec.name], spec.name)
            for spec in DATA_SHEETS
            if spec.name in wb.sheetnames
        }
        scan = _scan_info(wb["Information"])
    finally:
        wb.close()

    findings: dict[str, list[Finding]] = defaultdict(list)
    for r in sheets["Findings"]:
        findings[_text(r.get("Asset key"))].append(
            Finding(
                feature_id=_text(r.get("Rule")),
                category=_text(r.get("Category")),
                occurrences=int(_num(r.get("Occurrences")) or 1),
                penalty=int(_num(r.get("Penalty"))),
                support=_text(r.get("Fabric support")) or "unknown",
                guidance=_text(r.get("What to do in Fabric")),
            )
        )
    metrics = _metrics(sheets)
    assets: list[AssetInput] = []
    seen: set[str] = set()
    for r in sheets["Assets"]:
        key = _text(r.get("Asset key")) or f"{_text(r.get('Asset type'))}:{_text(r.get('Name'))}"
        if key in seen:
            raise InvalidWorkbookError(f"{path.name}: asset key {key} appears twice on Assets.")
        seen.add(key)
        assets.append(
            AssetInput(
                key=key,
                asset_type=_text(r.get("Asset type")),
                name=_text(r.get("Name")),
                location=_text(r.get("Location / id")),
                fabric_target=_text(r.get("Fabric target")),
                readiness_score=int(_num(r.get("Readiness score"))),
                readiness=_text(r.get("Readiness")) or "Compatible",
                findings=tuple(findings.get(key, ())),
                metrics=metrics.get(key, {}),
            )
        )
    return DiscoveryWorkbook(
        file_name=path.name, scan=scan, assets=tuple(assets), counts=item_counts(sheets)
    )
