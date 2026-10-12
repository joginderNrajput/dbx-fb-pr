from __future__ import annotations

import itertools
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from dbx_premig.config import Settings, load_settings
from dbx_premig.estimation import summarise
from dbx_premig.excel.discovery_writer import build_workbook, write_discovery_workbook
from dbx_premig.excel.layout import DATA_START_ROW, HEADER_ROW
from dbx_premig.excel.writer import CHART_SHEETS, DATA_SHEETS, SHEET_ORDER, column_letter
from dbx_premig.models import DiscoveryReport, WarningRow

ROOT = Path(__file__).parent.parent
_TOTALS_REF = re.compile(r"'Totals'!\$([A-Z]+)\$(\d+)")


@pytest.fixture
def wb(report: DiscoveryReport, settings: Settings) -> Workbook:
    return build_workbook(report, settings)


def test_sheet_order(wb: Workbook) -> None:
    assert tuple(wb.sheetnames) == SHEET_ORDER
    assert all(wb[n].sheet_state == "hidden" for n in CHART_SHEETS)
    assert wb.active.title == "Discovery"


def test_data_sheets_headers_and_rows(wb: Workbook, report: DiscoveryReport) -> None:
    for spec in DATA_SHEETS:
        ws = wb[spec.name]
        headers = [ws.cell(row=HEADER_ROW, column=i + 1).value for i in range(len(spec.columns))]
        assert headers == [c.header for c in spec.columns], spec.name
        assert ws.freeze_panes == f"A{DATA_START_ROW}"
        n = len(spec.rows(report))
        if n:
            assert ws.cell(row=DATA_START_ROW + n - 1, column=1).value is not None, spec.name
            assert ws.cell(row=DATA_START_ROW + n, column=1).value is None, spec.name


def test_numeric_columns_right_aligned(wb: Workbook, report: DiscoveryReport) -> None:
    for spec in DATA_SHEETS:
        if not spec.rows(report):
            continue
        ws = wb[spec.name]
        for i, col in enumerate(spec.columns, start=1):
            if col.numeric:
                assert ws.cell(row=DATA_START_ROW, column=i).alignment.horizontal == "right"


def test_dashboard_is_only_formulas_into_totals(wb: Workbook) -> None:
    """Every number on Discovery and the chart helpers is a reference to a labelled Totals row."""
    totals = wb["Totals"]
    refs = 0
    for name in ("Discovery", *CHART_SHEETS):
        for row in wb[name].iter_rows():
            for cell in row:
                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith("=")
                    and "Totals" in cell.value
                ):
                    for col, r in _TOTALS_REF.findall(cell.value):
                        refs += 1
                        label_col = "A"
                        assert totals[f"{label_col}{r}"].value, f"{name}!{cell.coordinate}"
                        assert totals[f"{col}{r}"].value is not None
    assert refs >= 40
    # Tiles hold no literal numbers.
    for row in wb["Discovery"].iter_rows(min_row=7, max_row=26, min_col=8, max_col=42):
        for cell in row:
            assert not isinstance(cell.value, int | float)


def test_totals_are_live_formulas(wb: Workbook) -> None:
    t = wb["Totals"]
    values = []
    for r in range(7, t.max_row + 1):
        label, value = t.cell(row=r, column=1).value, t.cell(row=r, column=2).value
        if label == "Asset type":  # the metric sections end here
            break
        if label not in (None, "Metric") and value is not None:
            values.append(value)
    assert len(values) > 60
    assert all(isinstance(v, str) and v.startswith("=") for v in values)


def test_formula_columns_stay_put() -> None:
    """Catalogs / Schemas / Jobs formulas address these columns by letter."""
    by = {s.name: s for s in DATA_SHEETS}
    assert column_letter(by["Tables"], "Catalog") == "A"
    assert column_letter(by["Tables"], "Schema") == "B"
    assert column_letter(by["Tables"], "Table type") == "F"
    assert column_letter(by["JobTasks"], "Job id") == "A"
    assert column_letter(by["Jobs"], "Job id") == "A"
    assert column_letter(by["Schemas"], "Catalog") == "A"
    for name in ("Volumes", "Functions"):
        assert column_letter(by[name], "Catalog") == "B"
        assert column_letter(by[name], "Schema") == "C"


def test_branding_on_every_visible_sheet(wb: Workbook) -> None:
    for name in SHEET_ORDER:
        if name in CHART_SHEETS:
            continue
        assert len(wb[name]._images) == 1, name  # one composed banner picture


def test_client_text_is_never_a_formula(report: DiscoveryReport, settings: Settings) -> None:
    evil = report.estate.model_copy(
        update={"warnings": (WarningRow(source='=HYPERLINK("http://x","y")', message="=1+1"),)}
    )
    wb = build_workbook(report.model_copy(update={"estate": evil}), settings)
    ws = wb["Warnings"]
    assert ws["A7"].data_type == "s" and ws["C7"].data_type == "s"


def test_long_text_truncated_and_control_chars_removed(
    report: DiscoveryReport,
) -> None:
    s = load_settings(max_cell_chars=100)
    est = report.estate.model_copy(
        update={"warnings": (WarningRow(source="a\x00b", message="x" * 500),)}
    )
    wb = build_workbook(report.model_copy(update={"estate": est}), s)
    assert wb["Warnings"]["A7"].value == "ab"
    assert len(wb["Warnings"]["C7"].value) == 100


def test_charts_do_not_overlap(wb: Workbook) -> None:
    from dbx_premig.excel.charts import _SIZE, CHART_ANCHORS

    ws = wb["Discovery"]
    assert len(ws._charts) == 3
    cols = [ws[a].column for a in CHART_ANCHORS]
    width_cols = _SIZE[0] / 2.54 * 96 / 28  # cm -> px -> 3.3-wide columns (~28 px)
    assert all(b - a >= width_cols for a, b in itertools.pairwise(cols))


def test_write_is_atomic_and_errors_are_clean(
    tmp_path: Path, report: DiscoveryReport, settings: Settings
) -> None:
    out = write_discovery_workbook(report, tmp_path / "d.xlsx", settings)
    assert load_workbook(out).sheetnames[0] == "Discovery"
    assert [p.name for p in tmp_path.iterdir()] == ["d.xlsx"]
    from dbx_premig.errors import ReportWriteError

    blocker = tmp_path / "file"
    blocker.write_text("x")
    with pytest.raises(ReportWriteError):
        write_discovery_workbook(report, blocker / "d.xlsx", settings)


def test_no_hex_colours_outside_theme() -> None:
    hex_re = re.compile(r"['\"](?:[0-9A-Fa-f]{6}|[0-9A-Fa-f]{8})['\"]")
    offenders = [
        str(p.relative_to(ROOT))
        for p in (ROOT / "dbx_premig" / "excel").rglob("*.py")
        if p.name != "theme.py" and hex_re.search(p.read_text())
    ]
    assert not offenders


@pytest.mark.slow
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_totals_recalculate_to_python_summary(
    tmp_path: Path, report: DiscoveryReport, settings: Settings
) -> None:
    """LibreOffice recalculates every formula; the results must equal the Python roll-up."""
    src = write_discovery_workbook(report, tmp_path / "in.xlsx", settings)
    outdir = tmp_path / "calc"
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True,
        capture_output=True,
        timeout=240,
    )
    wb = load_workbook(outdir / "in.xlsx", data_only=True)
    t = wb["Totals"]
    got = {
        t.cell(row=r, column=1).value: t.cell(row=r, column=2).value
        for r in range(7, t.max_row + 1)
    }
    s = summarise(report)
    assert got["Assessable objects"] == s.total_assets
    for level, n in s.by_readiness.items():
        assert got[f"{level} assets"] == n
    assert got["Notebooks and code files"] == len(report.estate.notebooks)
    assert got["Columns"] == len(report.estate.columns)
    assert got["Discovery warnings"] == s.warnings
    errors = [
        c.coordinate
        for ws in wb
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith(("#REF", "#VALUE", "#NAME", "#DIV"))
    ]
    assert not errors


def _banner_ok(ws: Any) -> None:
    from openpyxl.utils.units import EMU_to_pixels

    from dbx_premig.excel.banner import BAND_HEIGHT_PX

    (img,) = ws._images
    anchor = img.anchor
    assert type(anchor).__name__ == "OneCellAnchor", ws.title
    assert (anchor._from.col, anchor._from.colOff, anchor._from.row, anchor._from.rowOff) == (
        0, 0, 0, 0)  # fmt: skip
    assert EMU_to_pixels(anchor.ext.height) == BAND_HEIGHT_PX
    assert EMU_to_pixels(anchor.ext.width) >= 1500


def test_banner_is_one_picture_at_a1(
    wb: Workbook, report: DiscoveryReport, settings: Settings, tmp_path: Path
) -> None:
    """Logo, title and platform logo are composed into ONE picture at A1 on every sheet, so
    Excel cannot shift them onto each other (separate column-anchored pictures could)."""
    from dbx_premig.service import run_assessment

    discovery = write_discovery_workbook(report, tmp_path / "d.xlsx", settings)
    assessed = load_workbook(
        run_assessment(settings, discovery_workbook=discovery, output_dir=tmp_path).workbook
    )
    for book in (wb, assessed):
        for ws in book.worksheets:
            if ws.sheet_state != "hidden":
                _banner_ok(ws)


def test_compose_banner_sizes_and_long_titles() -> None:
    import io

    from PIL import Image

    from dbx_premig.excel.banner import BAND_HEIGHT_PX, compose_banner

    img = Image.open(io.BytesIO(compose_banner("Jobs  ·  Databricks Workflows", 1600)))
    assert img.size == (3200, BAND_HEIGHT_PX * 2) and img.format == "JPEG"
    long = Image.open(io.BytesIO(compose_banner("X" * 400, 1500)))
    assert long.size == (3000, BAND_HEIGHT_PX * 2)
    assert compose_banner("a", 1500) == compose_banner("a", 1500)  # deterministic
