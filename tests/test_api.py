from __future__ import annotations

import time
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr

from dbx_premig.api.app import create_app
from dbx_premig.config import Settings, load_settings
from dbx_premig.service import run_discovery
from tests.conftest import SNAPSHOT
from tests.fake_workspace import TOKEN, FakeWorkspace, serve


@pytest.fixture
def discovery_xlsx(tmp_path: Path) -> Path:
    return run_discovery(
        load_settings(), input_path=SNAPSHOT, output_dir=tmp_path / "d", output_name="d.xlsx"
    ).workbook


def _client(tmp_path: Path, **kw: Any) -> TestClient:
    settings = load_settings(api_work_dir=tmp_path / "runs", retry_backoff_s=0.0, **kw)
    return TestClient(create_app(settings))


def _wait(client: TestClient, url: str) -> dict[str, Any]:
    for _ in range(200):
        body: dict[str, Any] = client.get(url).json()
        if body["status"] in {"succeeded", "failed"}:
            return body
        time.sleep(0.05)
    raise AssertionError("run did not finish")


def test_healthz(tmp_path: Path) -> None:
    r = _client(tmp_path).get("/healthz")
    assert r.status_code == 200 and r.json()["database"] == "not configured"


def test_upload_workbook_runs_assessment_only(tmp_path: Path, discovery_xlsx: Path) -> None:
    client = _client(tmp_path)
    with discovery_xlsx.open("rb") as fh:
        r = client.post(
            "/v1/assessments",
            data={"organizationId": "org", "jobId": "job-1"},
            files={"file": ("d.xlsx", fh, "application/octet-stream")},
        )
    assert r.status_code == 202, r.text
    body = _wait(client, r.json()["links"]["status"])
    assert body["status"] == "succeeded", body
    assert body["mode"] == "workbook" and body["assessment"]["total_assets"] == 41
    assert body["persisted"] is None
    xlsx = client.get(body["links"]["assessmentWorkbook"])
    assert xlsx.status_code == 200
    out = tmp_path / "a.xlsx"
    out.write_bytes(xlsx.content)
    assert load_workbook(out).sheetnames[0] == "Assessment"
    assert client.get(body["links"]["discoveryWorkbook"]).status_code == 200


def test_databricks_json_runs_discovery_then_assessment(tmp_path: Path) -> None:
    client = _client(tmp_path)
    with serve(FakeWorkspace.from_fixture()) as url:
        r = client.post(
            "/v1/assessments/databricks",
            json={
                "organizationId": "org",
                "jobId": "job-2",
                "workspaceId": url,
                "token": TOKEN,
                "databricksCatalog": "main",
                "databricksSchema": "sales",
                "wait": True,
            },
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "succeeded" and body["mode"] == "databricks"
    assert body["discovery"]["counts"]["catalogs"] == 1  # scoped to catalog main
    assert body["discovery"]["counts"]["schemas"] == 1  # and schema sales
    assert TOKEN not in r.text


def test_multipart_credentials_and_oauth(tmp_path: Path) -> None:
    ws = FakeWorkspace.from_fixture()
    ws.oauth = ("sp", "secret")
    client = _client(tmp_path)
    with serve(ws) as url:
        r = client.post(
            "/v1/assessments",
            data={
                "organizationId": "o",
                "jobId": "j",
                "workspaceUrl": url,
                "clientId": "sp",
                "clientSecret": "secret",
                "wait": "true",
            },
        )
    assert r.status_code == 200, r.text
    assert r.json()["assessment"]["total_assets"] == 41


def test_bad_credentials_fail_cleanly(tmp_path: Path) -> None:
    client = _client(tmp_path)
    with serve(FakeWorkspace.from_fixture()) as url:
        r = client.post("/v1/assessments/databricks", json={
            "organizationId": "o", "jobId": "j", "workspaceUrl": url, "token": "nope",
            "wait": True})  # fmt: skip
    assert r.status_code == 422
    assert "401" in r.json()["error"] and "nope" not in r.text


@pytest.mark.parametrize(
    ("name", "content", "status"),
    [
        ("x.csv", b"a,b", 415),
        ("x.xlsx", b"not a zip", 415),
        ("x.xlsx", b"PK\x03\x04" + b"0" * 5000, 413),
    ],
)
def test_upload_rejected(tmp_path: Path, name: str, content: bytes, status: int) -> None:
    client = _client(tmp_path, max_upload_bytes=2048)
    r = client.post("/v1/assessments", data={"organizationId": "o", "jobId": name},
                    files={"file": (name, content)})  # fmt: skip
    assert r.status_code == status


def test_garbage_xlsx_fails_run(tmp_path: Path) -> None:
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "x")
    client = _client(tmp_path)
    r = client.post("/v1/assessments", data={"organizationId": "o", "jobId": "j", "wait": "1"},
                    files={"file": ("x.xlsx", buf.getvalue())})  # fmt: skip
    assert r.status_code == 422 and r.json()["status"] == "failed"


def test_validation_errors(tmp_path: Path) -> None:
    client = _client(tmp_path)
    assert (
        client.post("/v1/assessments", data={"organizationId": "o", "jobId": "j"}).status_code
        == 422
    )
    r = client.post("/v1/assessments", data={"organizationId": "o", "jobId": "j",
                                             "workspaceUrl": "https://x"})  # fmt: skip
    assert r.status_code == 422  # no token / client secret
    assert client.post("/v1/assessments/databricks", json={"jobId": "j"}).status_code == 422
    r = client.post("/v1/assessments", data={"organizationId": "o", "jobId": "j",
                                             "persist": "true", "workspaceUrl": "https://x",
                                             "token": "t"})  # fmt: skip
    assert r.status_code == 503  # persistence requested, no database


def test_conflict_and_unknown_and_expired(tmp_path: Path, discovery_xlsx: Path) -> None:
    app = create_app(load_settings(api_work_dir=tmp_path / "runs", api_run_ttl_s=60))
    client = TestClient(app)
    store = app.state.store
    run = store.create("o", "busy", "workbook")
    with discovery_xlsx.open("rb") as fh:
        r = client.post("/v1/assessments", data={"organizationId": "o", "jobId": "busy"},
                        files={"file": ("d.xlsx", fh)})  # fmt: skip
    assert r.status_code == 409
    assert client.get("/v1/runs/nope").status_code == 404
    store.finish(run, "succeeded")
    run.finished_mono = -10_000.0
    assert client.get(f"/v1/runs/{run.run_id}").status_code == 410
    assert not run.folder.exists()


def test_api_key(tmp_path: Path) -> None:
    client = _client(tmp_path, api_key=SecretStr("k"))
    assert client.get("/healthz").status_code == 200  # never needs a key
    assert client.get("/v1/runs/x").status_code == 401
    assert client.get("/v1/runs/x", headers={"X-API-Key": "k"}).status_code == 404


@pytest.mark.database
def test_api_persists_to_postgres(
    tmp_path: Path, discovery_xlsx: Path, pg_settings: Settings
) -> None:
    client = TestClient(create_app(pg_settings.model_copy(update={"api_work_dir": tmp_path})))
    with discovery_xlsx.open("rb") as fh:
        r = client.post("/v1/assessments", data={"organizationId": "api-org", "jobId": "api-job",
                                                 "wait": "true"},
                        files={"file": ("d.xlsx", fh)})  # fmt: skip
    assert r.status_code == 200, r.text
    persisted = r.json()["persisted"]
    assert persisted["rows"] == {"totals": 1, "inventory": 41, "finding": 70, "estimation": 41,
                                 "computation": 1}  # fmt: skip
