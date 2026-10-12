from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from typer.testing import CliRunner

from dbx_premig import cli
from dbx_premig.config import load_settings
from dbx_premig.estimation import summarise
from dbx_premig.service import build_report, run_discovery
from tests.conftest import FIXED_CLOCK, NOTEBOOKS, SNAPSHOT
from tests.fake_workspace import TOKEN, FakeWorkspace, serve

runner = CliRunner()


def _invoke(*args: str):  # type: ignore[no-untyped-def]
    return runner.invoke(cli.app, list(args))


def test_version_and_catalog_info() -> None:
    assert _invoke("version").stdout.startswith("dbx-premig ")
    r = _invoke("catalog-info")
    assert r.exit_code == 0 and "features" in r.stdout
    assert "notebooks" in _invoke("collectors").stdout


def test_discover_offline_stdout_is_only_the_path(tmp_path: Path) -> None:
    r = _invoke("discover", "--input", str(NOTEBOOKS), "--out", str(tmp_path), "--name", "x")
    assert r.exit_code == 0, r.output
    assert r.stdout.strip().splitlines()[-1] == str(tmp_path / "x.xlsx")
    assert (tmp_path / "x.xlsx").is_file()


def test_discover_live_against_fake_workspace(tmp_path: Path) -> None:
    with serve(FakeWorkspace.from_fixture()) as url:
        r = _invoke(
            "discover",
            "--host",
            url,
            "--token",
            TOKEN,
            "--out",
            str(tmp_path),
            "--name",
            "live.xlsx",
            "--save-snapshot",
            "--only",
            "jobs,clusters",
        )
    assert r.exit_code == 0, r.output
    snap = json.loads((tmp_path / "live_snapshot.json").read_text())
    assert set(snap["sections"]) == {"jobs", "clusters"}
    assert TOKEN not in (tmp_path / "live_snapshot.json").read_text()


@pytest.mark.parametrize(
    ("args", "code"),
    [
        (["--input", "does-not-exist"], 7),
        (
            [
                "--input",
                str(SNAPSHOT.parent / "notebooks" / "Shared" / "etl" / "report.sql"),
                "--only",
                "nope",
            ],
            11,
        ),
        ([], 11),  # live without host / token
    ],
)
def test_exit_codes(tmp_path: Path, args: list[str], code: int) -> None:
    r = _invoke("discover", "--out", str(tmp_path), *args)
    assert r.exit_code == code, r.output
    assert "Traceback" not in r.output


def test_unsupported_and_invalid_snapshot(tmp_path: Path) -> None:
    bad = tmp_path / "x.txt"
    bad.write_text("x")
    assert _invoke("discover", "--input", str(bad), "--out", str(tmp_path)).exit_code == 2
    snap = tmp_path / "s.json"
    snap.write_text("{}")
    assert _invoke("discover", "--input", str(snap), "--out", str(tmp_path)).exit_code == 8


def test_401_exit_code(tmp_path: Path) -> None:
    with serve(FakeWorkspace.from_fixture()) as url:
        r = _invoke("discover", "--host", url, "--token", "wrong", "--out", str(tmp_path))
    assert r.exit_code == 10
    assert "wrong" not in r.output


def test_internal_error_is_exit_1_without_traceback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(*_a: object, **_k: object) -> None:
        raise RuntimeError("secret detail")

    monkeypatch.setattr("dbx_premig.service.run_discovery", boom)
    r = _invoke("discover", "--input", str(NOTEBOOKS), "--out", str(tmp_path))
    assert r.exit_code == 1 and "RuntimeError" in r.output and "secret detail" not in r.output
    r = _invoke("discover", "--input", str(NOTEBOOKS), "--out", str(tmp_path), "-v")
    assert r.exit_code == 1 and "Traceback" in r.output


def test_cli_and_service_parity(tmp_path: Path) -> None:
    """The CLI adds no logic: its report equals the service's, field by field."""
    s = load_settings()
    direct = run_discovery(s, input_path=SNAPSHOT, output_dir=tmp_path / "a", output_name="a")
    r = _invoke("discover", "--input", str(SNAPSHOT), "--out", str(tmp_path / "b"), "--name", "b")
    assert r.exit_code == 0
    from openpyxl import load_workbook

    a = load_workbook(direct.workbook)["Assets"]
    b = load_workbook(tmp_path / "b" / "b.xlsx")["Assets"]
    rows_a = [[c.value for c in row] for row in a.iter_rows(min_row=6)]
    rows_b = [[c.value for c in row] for row in b.iter_rows(min_row=6)]
    assert rows_a == rows_b


def test_report_deterministic() -> None:
    s = load_settings()
    from dbx_premig.reader import read_snapshot

    snap = read_snapshot(SNAPSHOT, s)
    a = build_report(snap, s, generated_at=FIXED_CLOCK)
    b = build_report(snap, s, generated_at=FIXED_CLOCK)
    assert a == b and summarise(a) == summarise(b)


def test_interactive_without_tty_reads_input_from_stdin(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    from dbx_premig import prompts

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "stdin", io.StringIO(f"{SNAPSHOT}\n"))
    assert prompts.run_interactive() == 0
    out = capsys.readouterr().out.strip().splitlines()[-1]
    assert out.endswith(".xlsx") and (tmp_path / out).is_file()


def test_interactive_error_exit_code(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from dbx_premig import prompts

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "stdin", io.StringIO("missing-folder\n"))
    assert prompts.run_interactive() == 7


def test_main_without_args_runs_prompts(monkeypatch: pytest.MonkeyPatch) -> None:
    from dbx_premig import prompts

    monkeypatch.setattr(sys, "argv", ["discovery"])
    monkeypatch.setattr(prompts, "run_interactive", lambda: 0)
    with pytest.raises(SystemExit) as exc:
        cli.main()
    assert exc.value.code == 0
