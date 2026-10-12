from __future__ import annotations

import shutil
import subprocess
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dbx_premig.config import Settings, load_settings
from dbx_premig.errors import InvalidWorkbookError, RulesError
from dbx_premig.estimation.effort import assess_workbook, estimate_asset
from dbx_premig.estimation.formula import Formula, FormulaError
from dbx_premig.estimation.rules import load_rules, parse_rules
from dbx_premig.excel.assessment_writer import ASSESSMENT_ORDER, build_assessment_workbook
from dbx_premig.excel.discovery_reader import read_discovery_workbook
from dbx_premig.models import AssetInput, DiscoveryReport, Finding
from dbx_premig.resources import rules_path
from dbx_premig.service import run_assessment, run_discovery
from tests.conftest import FIXED_CLOCK, SNAPSHOT

RULES_TEXT = rules_path().read_text()


@pytest.fixture
def discovery_xlsx(tmp_path: Path, settings: Settings) -> Path:
    return run_discovery(settings, input_path=SNAPSHOT, output_dir=tmp_path,
                         output_name="d.xlsx").workbook  # fmt: skip


# --------------------------------------------------------------------------- formula


@pytest.mark.parametrize(
    ("expr", "env", "expected"),
    [
        ("1 + 2 * 3", {}, 7),
        ("lines / 40 + cells", {"lines": 80, "cells": 1}, 3.0),
        ("min(100, round(x))", {"x": 250.4}, 100),
        ("a if a > b else b", {"a": 1, "b": 2}, 2),
        ("readiness == 'Manual' or score >= 50", {"readiness": "Partial", "score": 50}, True),
        ("not true", {}, False),
        ("round(2.5) + round(0.125, 2)", {}, 3.13),
        ("x / 0", {"x": 5}, 0),
        ("ceil(1.2) + floor(1.8) + abs(-2) + sqrt(4)", {}, 7.0),
        ("1 < x <= 3", {"x": 3}, True),
    ],
)
def test_formula_evaluates(expr: str, env: dict[str, object], expected: object) -> None:
    assert Formula.parse(expr).evaluate(env) == expected  # type: ignore[arg-type]


@pytest.mark.parametrize(
    "expr",
    [
        "__import__('os')",
        "x.__class__",
        "open('f')",
        "[1, 2]",
        "lambda: 1",
        "x[0]",
        "2 ** 1000",
        "",
        "1 +",
        "max(key=1)",
        "{1: 2}",
        "None",
    ],
)
def test_formula_rejects_unsafe_or_invalid(expr: str) -> None:
    with pytest.raises(FormulaError):
        Formula.parse(expr).evaluate({"x": 1})


def test_formula_unknown_names_caught_at_parse() -> None:
    with pytest.raises(FormulaError, match="lnes"):
        Formula.parse("lnes / 10", frozenset({"lines"}))
    with pytest.raises(FormulaError, match="must give a number"):
        Formula.parse("'a'").number({})


# --------------------------------------------------------------------------- rules


def test_bundled_rules_load() -> None:
    rules = load_rules()
    assert rules.version == 1 and "default" in rules.base_hours
    assert rules.label.startswith("assessment_rules.yaml v1 (")


@pytest.mark.parametrize(
    ("old", "new", "match"),
    [
        ('score: "min(100, round(penalty_total + size_points))"', 'score: "bogus * 2"', "bogus"),
        ("    - level: Simple\n", "    - level: Easy\n", "exactly the levels"),
        ("  partial: 0.5\n", "  partial: 2\n", "between 0 and 1"),
        ('    default: "0"\n    Notebook: "lines / 500"', '    Notebook: "lines / 500"', "default"),
        ("version: 1\n", "version: 2\n", "version must be 1"),
        ("    Manual:              { dev: 1.0,  test: 1.0 }\n", "", "automation_factors"),
        (
            '  manual_dev_hours: "base_hours + size_hours"\n',
            '  manual_dev_hours: "manual_test_hours"\n',
            "manual_test_hours",
        ),
        ("hours_decimal_places: 2", "hours_decimal_places: 9", "0-6"),
    ],
)
def test_invalid_rules_rejected(old: str, new: str, match: str) -> None:
    assert old in RULES_TEXT, old
    with pytest.raises(RulesError, match=match):
        parse_rules(RULES_TEXT.replace(old, new), "custom.yaml")


def test_rules_file_errors(tmp_path: Path) -> None:
    with pytest.raises(RulesError):
        load_rules(tmp_path / "missing.yaml")
    with pytest.raises(RulesError):
        parse_rules("a: [", "x.yaml")
    with pytest.raises(RulesError):
        parse_rules("- 1", "x.yaml")


def test_custom_rules_change_the_numbers(discovery_xlsx: Path, settings: Settings) -> None:
    dw = read_discovery_workbook(discovery_xlsx, settings)
    base = assess_workbook(dw, load_rules(), generated_at=FIXED_CLOCK).computation
    doubled = parse_rules(
        RULES_TEXT.replace('manual_dev_hours: "base_hours + size_hours"',
                           'manual_dev_hours: "(base_hours + size_hours) * 2"'),
        "double.yaml",
    )  # fmt: skip
    other = assess_workbook(dw, doubled, generated_at=FIXED_CLOCK).computation
    assert other.manual_dev_hours == pytest.approx(base.manual_dev_hours * 2, abs=0.05)
    assert other.rules_version.startswith("double.yaml v1")


# --------------------------------------------------------------------------- engine


def _asset(**kw: object) -> AssetInput:
    data: dict[str, object] = {"key": "Notebook:/a", "asset_type": "Notebook", "name": "a"}
    data.update(kw)
    return AssetInput.model_validate(data)


def test_estimate_by_hand() -> None:
    rules = load_rules()
    f = Finding(feature_id="DBU-WIDGETS", category="dbutils", occurrences=1, penalty=6,
                support="partial", guidance="")  # fmt: skip
    a = _asset(readiness="Compatible", readiness_score=94, findings=(f,),
               metrics={"lines": 400, "cells": 8})  # fmt: skip
    e = estimate_asset(a, rules)
    # size = 400/40 + 8/4 = 12; score = round(6 + 12) = 18 -> Simple
    assert (e.complexity_score, e.complexity_level) == (18, "Simple")
    # constructs: 1 base + 1 finding; supported 1 + 0.5
    assert (e.total_constructs, e.supported, e.unsupported) == (2, 1.5, 0.5)
    assert e.automation_level == "Fully automated"
    # manual dev = 2 (Notebook Simple) + 400/500 = 2.8; test = 0.98
    assert (e.manual_dev_hours, e.manual_test_hours) == (2.8, 0.98)
    assert (e.automation_dev_hours, e.automation_test_hours) == (0.56, 0.49)
    assert e.hours_saved == round(2.8 + 0.98 - 0.56 - 0.49, 2)


def test_manual_and_unsupported_levels() -> None:
    rules = load_rules()
    assert estimate_asset(_asset(readiness="Manual", readiness_score=30), rules
                          ).automation_level == "Manual"  # fmt: skip
    e = estimate_asset(_asset(asset_type="Serving Endpoint", readiness="Unsupported"), rules)
    assert e.automation_level == "Manual" and e.hours_saved == 0


def test_assessment_of_fixture(discovery_xlsx: Path, settings: Settings) -> None:
    dw = read_discovery_workbook(discovery_xlsx, settings)
    r = assess_workbook(dw, load_rules(), generated_at=FIXED_CLOCK)
    c = r.computation
    assert c.total_assets == 41 == len(r.estimations)
    assert c.simple_count + c.medium_count + c.complex_count == 41
    assert c.fully_automated + c.partially_automated + c.manual == 41
    assert c.supported_constructs + c.unsupported_constructs == pytest.approx(c.total_constructs)
    # totals are sums of the rounded per-asset figures
    assert c.manual_dev_hours == pytest.approx(sum(e.manual_dev_hours for e in r.estimations))
    assert c.total_manual_hours == pytest.approx(c.manual_dev_hours + c.manual_test_hours)
    assert c.saving_pct == pytest.approx(c.hours_saved / c.total_manual_hours, abs=1e-6)
    assert c.assessment_generated_at == FIXED_CLOCK.isoformat()
    # deterministic
    assert assess_workbook(dw, load_rules(), generated_at=FIXED_CLOCK) == r


# --------------------------------------------------------------------------- reader


def test_reader_round_trips_the_discovery(
    discovery_xlsx: Path, settings: Settings, report: DiscoveryReport
) -> None:
    dw = read_discovery_workbook(discovery_xlsx, settings)
    assert [a.key for a in dw.assets] == [a.key for a in report.assessments]
    by = {a.key: a for a in report.assessments}
    for a in dw.assets:
        src = by[a.key]
        assert (a.readiness, a.readiness_score, a.fabric_target) == (
            src.readiness, src.readiness_score, src.fabric_target)  # fmt: skip
        assert [f.feature_id for f in a.findings] == [f.feature_id for f in src.findings]
    nb = next(a for a in dw.assets if a.key == "Notebook:/Shared/etl/ingest_orders")
    assert nb.metrics["lines"] == 25 and nb.metrics["cells"] == 5
    job = next(a for a in dw.assets if a.key == "Job:11")
    assert job.metrics["tasks"] == 3
    assert dw.counts["notebooks"] == 6 and dw.counts["columns"] == 12
    assert dw.counts["assessable_assets"] == 41 and dw.scan.mode == "live"
    assert dw.scan.metastore == "metastore_westeurope"


def test_editing_a_row_changes_the_assessment(
    discovery_xlsx: Path, settings: Settings, tmp_path: Path
) -> None:
    wb = load_workbook(discovery_xlsx)
    ws = wb["Assets"]
    for row in ws.iter_rows(min_row=7):
        if row[0].value == "Dashboard":
            row[7].value = "Manual"  # Readiness column
    # Editing a total must NOT change anything.
    wb["Totals"]["B9"] = 999
    edited = tmp_path / "edited.xlsx"
    wb.save(edited)
    dw = read_discovery_workbook(edited, settings)
    assert dw.counts["manual_assets"] == 4 and dw.counts["assessable_assets"] == 41
    r = assess_workbook(dw, load_rules(), generated_at=FIXED_CLOCK)
    assert sum(1 for e in r.estimations if e.asset_type == "Dashboard"
               and e.automation_level == "Manual") == 2  # fmt: skip


def test_reader_rejects_non_discovery_files(tmp_path: Path, settings: Settings) -> None:
    from openpyxl import Workbook

    other = tmp_path / "other.xlsx"
    Workbook().save(other)
    with pytest.raises(InvalidWorkbookError, match="not a discovery workbook"):
        read_discovery_workbook(other, settings)
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"nope")
    with pytest.raises(InvalidWorkbookError):
        read_discovery_workbook(junk, settings)
    with pytest.raises(InvalidWorkbookError):
        read_discovery_workbook(tmp_path / "d.csv", settings)
    wb = Workbook()
    wb.active.title = "Assets"
    wb.create_sheet("Findings")
    wb.create_sheet("Information")
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    with pytest.raises(InvalidWorkbookError, match="missing column"):
        read_discovery_workbook(bad, settings)


def test_reader_refuses_zip_bomb_workbook(discovery_xlsx: Path) -> None:
    with pytest.raises(InvalidWorkbookError, match="refused"):
        read_discovery_workbook(discovery_xlsx, load_settings(max_archive_ratio=1.5))


# --------------------------------------------------------------------------- workbook


def test_assessment_workbook_structure(
    discovery_xlsx: Path, settings: Settings, tmp_path: Path
) -> None:
    run = run_assessment(settings, discovery_workbook=discovery_xlsx, output_dir=tmp_path,
                         output_name="a", generated_at=FIXED_CLOCK)  # fmt: skip
    assert run.workbook.name == "a.xlsx"
    wb = load_workbook(run.workbook)
    assert tuple(wb.sheetnames) == ASSESSMENT_ORDER
    est = wb["Estimation"]
    assert est["O7"].value == "=M7+N7" and est["R7"].value == "=P7+Q7"
    summary = wb["Summary"]
    labels = {summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value
              for r in range(9, 27)}  # fmt: skip
    assert str(labels["Assets assessed"]).startswith("=COUNTA(")
    assert all(str(v).startswith("=") for v in labels.values() if v is not None)
    assert len(wb["Assessment"]._charts) == 3
    assert wb["Rules"]["B11"].value == RULES_TEXT.splitlines()[0]
    # dashboard tiles hold only formulas into Summary
    for row in wb["Assessment"].iter_rows(min_row=7, max_row=20, min_col=8, max_col=42):
        for cell in row:
            assert not isinstance(cell.value, int | float)
    default_name = run_assessment(
        settings, discovery_workbook=discovery_xlsx, output_dir=tmp_path, generated_at=FIXED_CLOCK
    ).workbook.name
    assert default_name == "databricks_assessment_d_20261001_1200.xlsx"


@pytest.mark.slow
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_assessment_summary_recalculates_to_computation(
    discovery_xlsx: Path, settings: Settings, tmp_path: Path
) -> None:
    run = run_assessment(settings, discovery_workbook=discovery_xlsx, output_dir=tmp_path,
                         output_name="a.xlsx")  # fmt: skip
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir",
                    str(tmp_path / "calc"), str(run.workbook)],
                   check=True, capture_output=True, timeout=240)  # fmt: skip
    wb = load_workbook(tmp_path / "calc" / "a.xlsx", data_only=True)
    s = wb["Summary"]
    got = {s.cell(row=r, column=1).value: s.cell(row=r, column=2).value for r in range(9, 27)}
    c = run.result.computation
    assert got["Assets assessed"] == c.total_assets
    assert got["Complex"] == c.complex_count and got["Manual"] == c.manual
    assert got["Total manual hours"] == pytest.approx(c.total_manual_hours, abs=0.01)
    assert got["Total automation hours"] == pytest.approx(c.total_automation_hours, abs=0.01)
    assert got["Hours saved"] == pytest.approx(c.hours_saved, abs=0.01)
    assert got["Saving %"] == pytest.approx(c.saving_pct, abs=1e-4)


def test_build_workbook_with_no_assets_is_rejected(tmp_path: Path, settings: Settings) -> None:
    from dbx_premig.errors import NothingFoundError
    from dbx_premig.excel.discovery_writer import write_discovery_workbook
    from dbx_premig.models import RawSnapshot
    from dbx_premig.service import build_report

    snap = RawSnapshot(mode="live", source="x", captured_at="now",
                       sections={"catalogs": [{"name": "c"}]})  # fmt: skip
    path = write_discovery_workbook(build_report(snap, settings), tmp_path / "e.xlsx", settings)
    with pytest.raises(NothingFoundError):
        run_assessment(settings, discovery_workbook=path, output_dir=tmp_path)
    _ = (build_assessment_workbook, datetime, UTC)
