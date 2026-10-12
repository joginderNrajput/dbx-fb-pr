"""Check a workbook written by the BUILT binary (used by build.sh / build.ps1 and CI).

    python packaging/verify_workbook.py output.xlsx [--min-assets N]

Fails (exit 1) if the sheet order, branding images or the Totals formulas are wrong.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from dbx_premig.excel.writer import SHEET_ORDER


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--min-assets", type=int, default=1)
    parser.add_argument("--kind", choices=["discovery", "assessment"], default="discovery")
    args = parser.parse_args()
    wb = load_workbook(args.workbook)
    if args.kind == "assessment":
        return _assessment(wb, args)
    problems: list[str] = []
    if tuple(wb.sheetnames) != SHEET_ORDER:
        problems.append(f"sheet order differs: {wb.sheetnames}")
    for name in ("Discovery", "Assets", "Totals"):
        if name in wb.sheetnames and len(wb[name]._images) != 1:
            problems.append(f"{name}: expected one banner picture")
    assets = sum(1 for row in wb["Assets"].iter_rows(min_row=7, max_col=1) if row[0].value)
    if assets < args.min_assets:
        problems.append(f"Assets has {assets} rows, expected >= {args.min_assets}")
    totals = wb["Totals"]
    if not str(totals["B9"].value or "").startswith("=COUNTA("):
        problems.append("Totals!B9 is not a live formula")
    for problem in problems:
        print(f"FAIL: {problem}", file=sys.stderr)
    if not problems:
        print(f"OK: {args.workbook.name} ({len(wb.sheetnames)} sheets, {assets} assets)")
    return 1 if problems else 0


def _assessment(wb: Workbook, args: argparse.Namespace) -> int:
    from dbx_premig.excel.assessment_writer import ASSESSMENT_ORDER

    problems = []
    if tuple(wb.sheetnames) != ASSESSMENT_ORDER:
        problems.append(f"sheet order differs: {wb.sheetnames}")
    rows = sum(1 for row in wb["Estimation"].iter_rows(min_row=7, max_col=1) if row[0].value)
    if rows < args.min_assets:
        problems.append(f"Estimation has {rows} rows, expected >= {args.min_assets}")
    if len(wb["Assessment"]._charts) != 3:
        problems.append("Assessment: charts missing")
    for problem in problems:
        print(f"FAIL: {problem}", file=sys.stderr)
    if not problems:
        print(f"OK: {args.workbook.name} ({rows} assets assessed)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
